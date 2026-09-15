#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""C7 计算链重算验证（crwu-audit-datacheck 能力型脚本，随技能安装）。

对 xlsx **可见区**公式做独立重算，与缓存值比对，输出客观差异清单；**只报差异、不下价值判断**。

纪律（防自造假阳性）：
- 只读**可见区**（H0）；引用到隐藏 sheet/行/列、外部工作簿、不支持的函数、解析失败 → 一律标
  「未重算(原因)」，**绝不判通过、绝不猜值**。
- 重算结果仅与缓存值做「是否一致」比对；不一致 = 差异（`C7-计算链`），一致 = 已核。引擎自身不判断
  谁对谁错。
- 支持：算术 `+ - * / ^`、括号、比较 `= <> < <= > >=`、`SUM/PRODUCT/ROUND/IF/MAX/MIN`、
  单元格/绝对引用 `$A$1`、跨表引用 `'表名'!A1`；字符串与 TRUE/FALSE 常量。
  其余（百分比 `%`、文本拼接 `&`、MOD 等未列函数）不实现 → 标「未重算」，绝不猜值。

用法：
    python3 scripts/recalc_check.py --workbook <raw.xlsx> [--out <diff.json>]
"""
from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


class Unavailable(Exception):
    """该格无法重算：携带原因（隐藏区/外部引用/不支持函数/解析失败/缺失）。"""


class EngineError(Unavailable):
    """引擎**自身**未能处理该格（非"规则上不支持"）。

    单独成类并在摘要里计数，是为了**不让代码缺陷伪装成"该格数据有问题"**：
    若兜底把所有异常都并进普通「未重算」，整表静默降级也看不出来。
    """


# 单元格值的类型／数值问题：Excel 自己也会报 `#VALUE!`／`#DIV/0!`，
# 属"该格不可重算"，不是引擎缺陷。例：`=A1-B1` 而 A1 是日期、B1 是文本。
_VALUE_ERRORS = (TypeError, ValueError, ArithmeticError, InvalidOperation)


# ---------------------------------------------------------------- tokenizer
_CELL = r"\$?[A-Z]{1,3}\$?\d{1,7}"
_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.\u4e00-\u9fff]*)"
_TOKEN = re.compile(
    r"(?P<ws>\s+)"
    r"|(?P<sheetref>" + _SHEET + r"!)(?=\s*\$?[A-Za-z]{1,3}\$?\d)"
    r"|(?P<cellref>" + _CELL + r")"
    r"|(?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)"
    r'|(?P<str>"[^"]*")'
    r"|(?P<func>[A-Za-z_][A-Za-z0-9_.]*)(?=\s*\()"
    r"|(?P<bool>TRUE|FALSE)"
    r"|(?P<op><>|<=|>=|<|>|=|[+\-*/^%&(),:])",
    re.VERBOSE | re.IGNORECASE,
)


def _tokenize(formula: str):
    tokens = []
    pos = 0
    while pos < len(formula):
        m = _TOKEN.match(formula, pos)
        if not m:
            raise Unavailable("解析失败：无法识别的片段 " + formula[pos:pos + 12].strip())
        kind = m.lastgroup
        if kind == "ws":
            pos = m.end()
            continue
        value = m.group(0)
        if kind == "num":
            tokens.append(("num", float(value)))
        elif kind == "str":
            tokens.append(("str", value[1:-1]))
        elif kind == "bool":
            tokens.append(("bool", value.upper() == "TRUE"))
        elif kind == "cellref":
            tokens.append(("cell", _parse_cell(value)))
        elif kind == "sheetref":
            tokens.append(("sheet", value[:-1].strip("'")))
        elif kind == "func":
            tokens.append(("func", value.upper()))
        elif kind == "op":
            if value == ":":
                tokens.append(("colon", None))
            else:
                tokens.append(("op", value))
        pos = m.end()
    return tokens


def _parse_cell(text: str):
    text = text.replace("$", "")
    i = 0
    while i < len(text) and not text[i].isdigit():
        i += 1
    return (text[:i], int(text[i:]))


# ---------------------------------------------------------------- evaluator
class Evaluator:
    def __init__(self, wbv, wbf, hidden):
        self.wbv = wbv          # data_only=True（缓存值）
        self.wbf = wbf          # data_only=False（公式串）
        self.hidden = hidden    # {hiddenSheets, hiddenRows, hiddenCols}
        self.sheet = None       # 当前工作表名

    # ---- 值获取（H0：隐藏区不可读 → Unavailable） ----
    def _cell_value(self, sheet: str, coord) -> float | str | bool:
        if sheet not in self.wbv.sheetnames:
            raise Unavailable(f"引用不存在的工作表 {sheet}")
        if sheet in self.hidden["hiddenSheets"]:
            raise Unavailable("引用隐藏工作表")
        col_letter, row = coord
        if row in set(self.hidden["hiddenRows"].get(sheet, [])):
            raise Unavailable("引用隐藏行")
        if col_letter in set(self.hidden["hiddenCols"].get(sheet, [])):
            raise Unavailable("引用隐藏列")
        v = self.wbv[sheet].cell(row, column_index_from_string(col_letter)).value
        if v is None:
            return 0.0
        return v

    def _resolve(self, token_idx, tokens):
        """token_idx 处为 cell 或 sheet 前缀；返回 (value, 下一个未消费下标)。支持 range。"""
        tok = tokens[token_idx]
        if tok[0] == "sheet":
            return self._resolve_ref(tokens, token_idx + 1, tok[1])
        return self._resolve_ref(tokens, token_idx, self.sheet)

    def _resolve_ref(self, tokens, i, sheet):
        coord = tokens[i][1]
        nxt = i + 1
        if nxt < len(tokens) and tokens[nxt][0] == "colon":
            end = tokens[nxt + 1][1]
            return self._range(sheet, coord, end), nxt + 2
        return self._cell_value(sheet, coord), nxt

    def _range(self, sheet, start, end):
        c1, r1 = column_index_from_string(start[0]), start[1]
        c2, r2 = column_index_from_string(end[0]), end[1]
        values = []
        for r in range(min(r1, r2), max(r1, r2) + 1):
            for c in range(min(c1, c2), max(c1, c2) + 1):
                values.append(self._cell_value(sheet, (get_column_letter(c), r)))
        return values

    # ---- 主入口 ----
    def eval(self, formula: str):
        """统一兜底：任何异常都必须变成「未重算(原因)」，**绝不让单格打断整表**。

        契约（本文件 docstring）要求"解析失败/除零/不支持 → 一律标未重算，绝不判通过、绝不猜值"。
        整改前只捕获 `Unavailable`，于是 `=A1-B1`（A1 日期、B1 文本）这类 `TypeError`
        会打穿整个 `run()`（实战：3-资产基础法.xlsx 直接崩，datacheck 整表不可用）。
        """
        try:
            return self._eval(formula)
        except Unavailable:
            raise
        except _VALUE_ERRORS as exc:
            raise Unavailable(f"类型不匹配或计算失败：{type(exc).__name__}: {exc}") from exc
        except Exception as exc:                    # 非预期 → 单独成类，摘要中计数暴露
            raise EngineError(f"引擎异常 {type(exc).__name__}: {exc}") from exc

    def _eval(self, formula: str):
        if re.search(r"\[\d*\]", formula):
            raise Unavailable("引用外部工作簿")
        formula = formula.lstrip()
        if formula.startswith("="):
            formula = formula[1:].lstrip()
        tokens = _tokenize(formula)
        value, nxt = self._expr(tokens, 0)
        if nxt != len(tokens):
            raise Unavailable("解析失败：存在未消费的片段")
        return value

    def _expr(self, tokens, i):
        return self._comparison(tokens, i)

    def _comparison(self, tokens, i):
        left, i = self._additive(tokens, i)
        while i < len(tokens) and tokens[i][0] == "op" and tokens[i][1] in ("=", "<>", "<", "<=", ">", ">="):
            op = tokens[i][1]
            right, i = self._additive(tokens, i + 1)
            left = _compare(op, left, right)
        return left, i

    def _additive(self, tokens, i):
        value, i = self._multiplicative(tokens, i)
        while i < len(tokens) and tokens[i][0] == "op" and tokens[i][1] in ("+", "-"):
            op = tokens[i][1]
            rhs, i = self._multiplicative(tokens, i + 1)
            value = _arith(op, value, rhs)
        return value, i

    def _multiplicative(self, tokens, i):
        value, i = self._unary(tokens, i)
        while i < len(tokens) and tokens[i][0] == "op" and tokens[i][1] in ("*", "/", "^"):
            op = tokens[i][1]
            rhs, i = self._unary(tokens, i + 1)
            value = _arith(op, value, rhs)
        return value, i

    def _unary(self, tokens, i):
        if i < len(tokens) and tokens[i][0] == "op" and tokens[i][1] == "-":
            v, i = self._unary(tokens, i + 1)
            return -v, i
        return self._primary(tokens, i)

    def _primary(self, tokens, i):
        tok = tokens[i]
        kind = tok[0]
        if kind == "num":
            return tok[1], i + 1
        if kind == "str":
            return tok[1], i + 1
        if kind == "bool":
            return tok[1], i + 1
        if kind == "cell":
            return self._resolve(i, tokens)
        if kind == "sheet":
            return self._resolve(i, tokens)
        if kind == "func":
            return self._function(tokens, i)
        if kind == "op" and tok[1] == "(":
            value, i = self._expr(tokens, i + 1)
            if i >= len(tokens) or tokens[i] != ("op", ")"):
                raise Unavailable("解析失败：缺右括号")
            return value, i + 1
        raise Unavailable("解析失败：不支持的语法 " + str(tok))

    def _function(self, tokens, i):
        name = tokens[i][1]
        j = i + 1
        if j >= len(tokens) or tokens[j] != ("op", "("):
            raise Unavailable("解析失败：函数缺左括号")
        # 收集参数（惰性：把每个参数包成闭包）
        args = []
        j += 1
        depth = 0
        start = j
        while j < len(tokens):
            t = tokens[j]
            if t[0] == "op" and t[1] == "(":
                depth += 1
            elif t[0] == "op" and t[1] == ")":
                if depth == 0:
                    if j > start:
                        args.append((start, j))
                    j += 1
                    break
                depth -= 1
            elif t[0] == "op" and t[1] == "," and depth == 0:
                args.append((start, j))
                j += 1
                start = j
                continue
            j += 1

        def ev(arg):
            s, e = arg
            return self._expr(tokens, s)[0]

        if name == "IF":
            if len(args) not in (2, 3):
                raise Unavailable("IF 参数个数错误")
            cond = ev(args[0])
            result = ev(args[1]) if _truthy(cond) else (ev(args[2]) if len(args) == 3 else False)
        elif name == "ROUND":
            if len(args) != 2:
                raise Unavailable("ROUND 参数个数错误")
            result = _round_half_away(ev(args[0]), int(ev(args[1])))
        elif name in ("SUM", "PRODUCT", "MAX", "MIN"):
            if not args:
                raise Unavailable(name + " 至少一个参数")
            values = [self._flatten(ev(a)) for a in args]
            flat = [v for grp in values for v in (grp if isinstance(grp, list) else [grp])]
            flat = [float(v) for v in flat]
            if name == "SUM":
                result = sum(flat)
            elif name == "PRODUCT":
                acc = 1.0
                for v in flat:
                    acc *= v
                result = acc
            elif name == "MAX":
                result = max(flat)
            else:
                result = min(flat)
        else:
            raise Unavailable("不支持的函数 " + name)
        return result, j

    @staticmethod
    def _flatten(v):
        return v if isinstance(v, list) else v


# ---------------------------------------------------------------- helpers
def _truthy(v):
    if isinstance(v, (int, float)):
        return v != 0
    if isinstance(v, str):
        return v != ""
    return bool(v)


def _compare(op, a, b):
    if op == "=":
        return a == b
    if op == "<>":
        return a != b
    if op == "<":
        return a < b
    if op == "<=":
        return a <= b
    if op == ">":
        return a > b
    return a >= b


def _arith(op, a, b):
    if op == "+":
        return a + b
    if op == "-":
        return a - b
    if op == "*":
        return a * b
    if op == "/":
        if b == 0:
            raise Unavailable("除零")
        return a / b
    if op == "^":
        return a ** b
    raise Unavailable("不支持的运算符 " + op)


def _round_half_away(x, n):
    # 用十进制字符串规避二进制浮点误差（如 ROUND(1.235,2) 应为 1.24，非 1.23）
    d = Decimal(str(x))
    q = Decimal(1).scaleb(-int(n))
    return float(d.quantize(q, rounding=ROUND_HALF_UP))


def _hidden_metadata(wbv):
    hidden = {"hiddenSheets": [], "hiddenRows": {}, "hiddenCols": {}}
    for ws in wbv.worksheets:
        if ws.sheet_state != "visible":
            hidden["hiddenSheets"].append(ws.title)
            continue
        rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        cols = {c for c, d in ws.column_dimensions.items() if d.hidden}
        if rows:
            hidden["hiddenRows"][ws.title] = sorted(rows)
        if cols:
            hidden["hiddenCols"][ws.title] = sorted(cols)
    return hidden


def _close(a, b) -> bool:
    if a == b:
        return True
    try:
        return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)), abs(float(b)))
    except (TypeError, ValueError):
        return False


# ---------------------------------------------------------------- main
def run(workbook: str):
    wbv = openpyxl.load_workbook(workbook, data_only=True)
    wbf = openpyxl.load_workbook(workbook, data_only=False)
    hidden = _hidden_metadata(wbv)
    ev = Evaluator(wbv, wbf, hidden)

    differences, not_recomputable = [], []
    formula_cells = recomputed = matched = mismatched = 0
    engine_errors = 0

    for ws in wbf.worksheets:
        if ws.sheet_state != "visible":
            continue
        hidden_rows = set(hidden["hiddenRows"].get(ws.title, []))
        hidden_cols = set(hidden["hiddenCols"].get(ws.title, []))
        ev.sheet = ws.title
        # 只遍历**实际存在**的格（精确集合），不用 `ws.iter_rows()`——后者按声明 dimension
        # 扫满矩形；原件常带游离的 `has_style=True`、`value=None` 空格（如整列刷格式残留），
        # 会把百万行全走一遍（2026-302135-LX9619-BG8634 实测：单表 1,048,575 行 → 长时间无响应）。
        # 与 crwu-audit 的 prepare_materials.py 同一口径。按 (行, 列) 排序以保持
        # 与原 `iter_rows()` 一致的行优序 —— 差异清单顺序必须可复现。
        for (row_i, col_i), cell in sorted(getattr(ws, "_cells", {}).items()):
            if not (isinstance(row_i, int) and isinstance(col_i, int)):
                continue
            if row_i in hidden_rows or get_column_letter(col_i) in hidden_cols:
                continue
            formula = cell.value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            formula_cells += 1
            cached = wbv[ws.title].cell(row_i, col_i).value
            try:
                recalculated = ev.eval(formula)
            except Unavailable as exc:
                entry = {"sheet": ws.title, "cell": cell.coordinate, "formula": formula,
                         "reason": str(exc)}
                if isinstance(exc, EngineError):
                    entry["engineError"] = True
                    engine_errors += 1
                not_recomputable.append(entry)
                continue
            except Exception as exc:                # 兜底：任何异常都不得中断整表
                engine_errors += 1
                not_recomputable.append(
                    {"sheet": ws.title, "cell": cell.coordinate, "formula": formula,
                     "engineError": True,
                     "reason": f"引擎异常 {type(exc).__name__}: {exc}"})
                continue
            recomputed += 1
            if cached is None or not _close(recalculated, cached):
                mismatched += 1
                differences.append({
                    "sheet": ws.title, "cell": cell.coordinate, "formula": formula,
                    "recalculated": recalculated, "cached": cached,
                    "note": "重算值与缓存值不一致" if cached is not None else "有公式但无缓存值",
                })
            else:
                matched += 1

    return {
        "file": workbook,
        "summary": {
            "formulaCells": formula_cells, "recomputed": recomputed,
            "matched": matched, "mismatched": mismatched,
            "notRecomputable": len(not_recomputable),
            # 引擎自身未能处理（类型不匹配等已归入普通未重算；此处只计非预期的引擎异常）。
            # 非 0 表示可能有引擎缺陷，须排查 —— 不得当作"该表数据问题"。
            "engineErrors": engine_errors,
        },
        "differences": differences,
        "notRecomputable": not_recomputable,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="C7 计算链重算验证（只报差异、不下判断）")
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    result = run(args.workbook)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            json.dump(result, fh, ensure_ascii=False, indent=1)
    print(json.dumps(result["summary"], ensure_ascii=False))
    for d in result["differences"]:
        print(f"  差异 {d['sheet']}!{d['cell']} 重算={d['recalculated']} 缓存={d['cached']} | {d['formula']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
