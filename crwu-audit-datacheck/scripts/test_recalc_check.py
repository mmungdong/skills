#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""`recalc_check.py` 契约测试（C7 计算链重算验证）。

自洽：只用 openpyxl 造临时 fixture，可随技能安装运行。每条对应真实/风险场景。
"""
from __future__ import annotations

import importlib.util
import json
import re
import tempfile
import unittest
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent


def _load():
    spec = importlib.util.spec_from_file_location("recalc_check_for_test",
                                                  SCRIPTS_DIR / "recalc_check.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _inject_cell(data, coord, val):
    """给指定公式格注入缓存值 `<v>`（字符串结果加 `t="str"`），模拟真实 Excel 已计算文件。"""
    # openpyxl 写公式格为 `<c r=".."><f>..</f><v /></c>`（空缓存）；把 `<v />` 替换为真实缓存值。
    pat = re.compile(rb'(<c r="%s"[^>]*?>)(<f>.*?</f>).*?(</c>)' % coord.encode())

    def repl(m):
        c = m.group(1)
        if isinstance(val, str):
            c = c[:-1] + b' t="str">'
        v = ("<v>%s</v>" % val).encode("utf-8")
        return c + m.group(2) + v + m.group(3)

    return pat.sub(repl, data, count=1)


def _inject_cached(path, cached):
    """`cached`: {(sheet_title, coord): value}。openpyxl 不写缓存值，故手术注入。"""
    import shutil
    import zipfile
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as z:
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    name_rid = {}
    for m in re.finditer(r"<sheet\b([^>]*?)/>", wbxml):
        attrs = m.group(1)
        nm = re.search(r'name="([^"]+)"', attrs)
        rid = re.search(r'r:id="([^"]+)"', attrs)
        if nm and rid:
            name_rid[nm.group(1)] = rid.group(1)
    rid_target = {}
    for m in re.finditer(r"<Relationship\b([^>]*?)/>", rels):
        attrs = m.group(1)
        rid = re.search(r'Id="([^"]+)"', attrs)
        tgt = re.search(r'Target="([^"]+)"', attrs)
        if rid and tgt:
            rid_target[rid.group(1)] = tgt.group(1)
    sheet_file = {n: rid_target[r].lstrip("/") for n, r in name_rid.items()}
    with zipfile.ZipFile(path) as z, zipfile.ZipFile(tmp, "w") as out:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename in sheet_file.values():
                sheet = next(n for n, f in sheet_file.items() if f == item.filename)
                for (s, coord), val in cached.items():
                    if s == sheet:
                        data = _inject_cell(data, coord, val)
            out.writestr(item, data)
    shutil.move(str(tmp), str(path))


class RecalcCheckContractTest(unittest.TestCase):
    def setUp(self):
        import openpyxl  # noqa: F401
        self.module = _load()
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, wb, name="t.xlsx"):
        import openpyxl
        p = self.dir / name
        wb.save(p)
        return self.module.run(str(p))

    def _run_cached(self, wb, cached, name="t.xlsx"):
        p = self.dir / name
        wb.save(p)
        _inject_cached(p, cached)
        return self.module.run(str(p))

    def test_arithmetic_and_absolute_ref_match(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["B1"], wb.active["C1"] = 3, 4, "=$A$1*B1"
        r = self._run_cached(wb, {("Sheet", "C1"): 12})
        self.assertEqual(1, r["summary"]["matched"], r["differences"])
        self.assertEqual([], r["differences"])

    def test_mismatch_is_reported(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["A2"] = 2, 3
        wb.active["A3"] = "=A1+A2"       # 公式=5，但缓存值故意写成 6
        wb.active["A3"].value = "=A1+A2"
        # 用 zip 手术注入缓存值 6
        import zipfile
        p = self.dir / "m.xlsx"
        wb.save(p)
        with zipfile.ZipFile(p) as z, zipfile.ZipFile(p.with_suffix(".tmp"), "w") as out:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet") and b"<f>A1+A2</f>" in data:
                    data = data.replace(b"<f>A1+A2</f>", b"<f>A1+A2</f><v>6</v>", 1)
                out.writestr(item, data)
        import shutil
        shutil.move(str(p.with_suffix(".tmp")), str(p))
        r = self.module.run(str(p))
        self.assertEqual(1, r["summary"]["mismatched"], r)
        self.assertEqual(5, r["differences"][0]["recalculated"])

    def test_sum_and_product_and_round(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["A2"], wb.active["A3"] = 2, 3, 4
        wb.active["A4"] = "=SUM(A1:A3)"
        wb.active["A5"] = "=PRODUCT(A1:A3)"
        wb.active["A6"] = "=ROUND(1.235,2)"
        r = self._run_cached(wb, {("Sheet", "A4"): 9, ("Sheet", "A5"): 24, ("Sheet", "A6"): 1.24})
        self.assertEqual(3, r["summary"]["matched"], r)
        self.assertEqual([], r["differences"])

    def test_if_chain_reproduces_real_case(self):
        """真实场景：描述词匹配 IF 链——匹配词错误时算 0（重算应复现该 0）。"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["D22"] = "较完善"
        ws["V22"] = '=IF(D22="完善",5,IF(D22="较完善",4,IF(D22="一般",3,0)))'  # 应为 4
        ws["V23"] = '=IF(D23="方便",5,IF(D23="较方便",4,0))'                     # D23 空 → 0
        r = self._run_cached(wb, {("Sheet", "V22"): 4, ("Sheet", "V23"): 0})
        self.assertEqual(2, r["summary"]["matched"], r)
        self.assertEqual([], r["differences"])

    def test_if_returns_correct_branch_and_lazy(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = '=IF(A1>5, "高", "低")'
        ws["A2"] = 3
        ws["B2"] = '=IF(A2>5, "高", "低")'
        r = self._run_cached(wb, {("Sheet", "B1"): "高", ("Sheet", "B2"): "低"})
        self.assertEqual(2, r["summary"]["matched"], r)

    def test_hidden_dependency_is_not_recomputable(self):
        """引用隐藏列 → 标「未重算」，绝不读隐藏值。"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "=N10*2"
        ws["N10"] = 999999
        ws.column_dimensions["N"].hidden = True
        r = self._run(wb)
        self.assertEqual(1, r["summary"]["notRecomputable"], r)
        self.assertIn("隐藏", r["notRecomputable"][0]["reason"])

    def test_cross_sheet_reference(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = 5
        s2 = wb.create_sheet("表2")
        s2["B2"] = "=Sheet!A1*2"   # 用工作表名（默认表名是 Sheet）
        r = self._run_cached(wb, {("表2", "B2"): 10})
        self.assertEqual(1, r["summary"]["matched"], r["differences"])

    def test_unsupported_function_is_not_recomputable(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = "=VLOOKUP(B1,C1:D5,2,FALSE)"
        r = self._run(wb)
        self.assertEqual(1, r["summary"]["notRecomputable"], r)
        self.assertIn("不支持的函数", r["notRecomputable"][0]["reason"])

    def test_external_workbook_reference_is_not_recomputable(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = "='[29]投资性房地产-房屋公允模式'!W31"
        r = self._run(wb)
        self.assertEqual(1, r["summary"]["notRecomputable"], r)
        self.assertIn("外部工作簿", r["notRecomputable"][0]["reason"])

    def test_division_by_zero_is_not_recomputable(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = "=1/0"
        r = self._run(wb)
        self.assertEqual(1, r["summary"]["notRecomputable"], r)
        self.assertIn("除零", r["notRecomputable"][0]["reason"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
