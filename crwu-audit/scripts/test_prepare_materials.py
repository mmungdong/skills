#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""`prepare_materials.py` 契约测试（H0 数据隔离 + 重建法三条纪律）。

自洽：只用本技能 `scripts/` 内的脚本造临时 fixture，可随技能一起安装运行。
每条用例都对应一次**真实失效**，防止回退：

1. 缓存值优先 —— 源格"公式 + 缓存值"时，重建后必须落**缓存值**（否则后续 data_only
   读回 None，把"未重算"读成"数据缺失"，曾误报"评估结果列全空"）。
2. H0 —— 人工隐藏的 sheet / 行 / 列 / 折叠分组整体排除：其内容值**不得出现在任何产物**里。
3. 同名消歧 —— 定稿与送审稿同名时，工作版必须两份都在、来源可区分、不得静默覆盖。
4. 「值不可得」登记 —— 无缓存值的公式格单独计数，不混进"数据缺失"。
5. 源材料目录只读 —— 运行后源目录文件集合不变。
6. 表格遍历四层边界 —— 扫描边界按**有值格**（不按 dimension、也不按"存在的格"）；
   声明用区远大于有值区时登记表格规范提示；单表超规模上限时记 capability gap 并跳过。
7. 媒体证据通道 —— 内嵌图/独立图片必须**导出到 `媒体证据/` 并写入 `媒体索引.json`**
   （此前只登记计数、无读取通道 → 证据在图片里被判"为空/缺失"）；阶段二复核件复跑
   必须换产物名，**不得覆盖阶段一冻结的 `材料盘点.json`**（复核意见附件里的图同样要被抽取）。
"""
from __future__ import annotations

import base64
import hashlib
import importlib.util
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent

_PNG = base64.b64decode(
    b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg==")


def _minimal_docx_with_image(path: Path) -> Path:
    """造一个正文第 2 段内联一张图的 `.docx`（复核意见附件常见形态）。"""
    ct = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
          'content-types"><Default Extension="rels" ContentType="application/vnd.'
          'openxmlformats-package.relationships+xml"/><Default Extension="xml" '
          'ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.'
          'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/'
            '2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>')
    doc = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
           'wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
           '2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
           'wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<w:body><w:p><w:r><w:t>问题 1：可比案例位置图缺失。</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>证据截图：</w:t></w:r><w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/>'
           '<wp:docPr id="1" name="P1"/><a:graphic><a:graphicData uri="http://schemas.'
           'openxmlformats.org/drawingml/2006/picture"><a:blip r:embed="rId10"/></a:graphicData>'
           '</a:graphic></wp:inline></w:drawing></w:r></w:p></w:body></w:document>')
    doc_rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/'
                '2006/relationships"><Relationship Id="rId10" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/image" Target="media/image1.png"/></Relationships>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", doc_rels)
        z.writestr("word/media/image1.png", _PNG)
    return path


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _load_module():
    spec = importlib.util.spec_from_file_location(
        "prepare_materials_for_test", SCRIPTS_DIR / "prepare_materials.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _with_cached_value(path: Path, cell_ref: str, formula: str, cached: str) -> None:
    """把 openpyxl 写出的公式格补上 Excel 才会写的 `<v>` 缓存值。

    直接造"公式 + 缓存值"的源格，才能复现线上那种"源有意义的值、重建后却读回 None"的场景。
    """
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                text = data.decode("utf-8")
                needle = f'<f>{formula}</f>'
                assert needle in text, f"{item.filename}: 未找到 {needle}"
                text = text.replace(needle, f"{needle}<v>{cached}</v>", 1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


class PrepareMaterialsContractTest(unittest.TestCase):
    def setUp(self):
        import openpyxl  # noqa: F401  （缺失则整体 skip）
        self.module = _load_module()
        self.tmp = tempfile.TemporaryDirectory()
        self.case = Path(self.tmp.name)
        self.src = self.case / "材料-源"
        self.src.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self):
        return self.module.prepare(
            str(self.case), str(self.src),
            str(self.case / "提取"), str(self.case / "工作版"),
        )

    def _payload(self):
        return json.loads((self.case / "材料盘点.json").read_text(encoding="utf-8"))

    def _inventory(self):
        return self._payload()["items"]

    def _run_label(self, label: str, src: Path, inventory: str = None,
                   media_index: str = None):
        """按阶段标签复跑（阶段二复核件用）：产物名加后缀，不影响阶段一冻结产物。"""
        return self.module.prepare(
            str(self.case), str(src),
            str(self.case / f"提取-{label}"), str(self.case / f"工作版-{label}"),
            extract_dir=str(self.case / f"解压-{label}"),
            media_dir=str(self.case / f"媒体证据-{label}"),
            inventory_name=inventory or f"{label}盘点.json",
            media_index_name=media_index or f"{label}媒体索引.json")

    # ---- 1 缓存值优先 ----------------------------------------------------
    def test_cached_value_wins_over_formula_string(self):
        """源格 = 公式 + 缓存值 → 重建后必须能 data_only 读回该值。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "明细表.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "评估结果"
        wb.active["A2"] = "=1+1"
        wb.save(f)
        _with_cached_value(f, "A2", "1+1", "42")

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/明细表.xlsx", data_only=True)
        self.assertEqual(42, rebuilt.active["A2"].value)
        self.assertEqual(
            0,
            self._inventory()[0]["workbook"]["sheets"][0]["valueUnavailable"],
            "有缓存值的格不应被记成「值不可得」",
        )

    def test_formula_without_cached_value_is_registered_as_value_unavailable(self):
        """无缓存值的公式格：保留公式串（不丢结构），并单独登记「值不可得」。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "=SUM(B1:B2)"
        wb.save(d / "外部引用.xlsx")

        self._run()

        rec = self._inventory()[0]
        self.assertEqual(1, rec["workbook"]["sheets"][0]["valueUnavailable"])
        self.assertTrue(rec["workbook"]["valueUnavailable"])
        rebuilt = openpyxl.load_workbook(self.case / "工作版/外部引用.xlsx", data_only=False)
        self.assertEqual("=SUM(B1:B2)", rebuilt.active["A1"].value)

    # ---- 2 H0 隐藏区 -----------------------------------------------------
    def test_hidden_content_never_reaches_any_artifact(self):
        """隐藏 sheet / 行 / 列的内容值不得出现在工作版或盘点中（H0）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "含隐藏.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "可见表"
        # A1/C1 可见（对照组），A2 隐藏行、B1 隐藏列、隐藏 sheet 各放一个哨兵值
        ws["A1"], ws["A2"], ws["B1"], ws["C1"] = "保留", 999999, 888888, 777777
        ws.row_dimensions[2].hidden = True           # 隐藏行
        ws.column_dimensions["B"].hidden = True      # 隐藏列
        hid = wb.create_sheet("隐藏表")
        hid["A1"] = 666666                           # 隐藏 sheet 的内容值
        hid.sheet_state = "hidden"                   # 必须显式设为隐藏（create_sheet 默认可见）
        wb.save(f)

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/含隐藏.xlsx", data_only=True)
        values = [c.value for row in rebuilt["可见表"].iter_rows() for c in row if c.value is not None]
        self.assertIn("保留", values)
        self.assertIn(777777, values, "可见列的值必须保留（对照组）")
        for sentinel in (999999, 888888, 666666):
            self.assertNotIn(sentinel, values, f"隐藏值 {sentinel} 泄漏到工作版")
        meta = self._inventory()[0]["workbook"]["hiddenMeta"]
        self.assertEqual(["隐藏表"], meta["hiddenSheets"], "隐藏 sheet 的**名字**必须登记（忽略清单）")
        self.assertEqual([2], meta["hiddenRows"]["可见表"], "隐藏行段位必须登记")
        self.assertEqual(["B"], meta["hiddenCols"]["可见表"], "隐藏列段位必须登记")
        blob = (self.case / "材料盘点.json").read_text(encoding="utf-8")
        for sentinel in ("999999", "888888", "666666"):
            self.assertNotIn(sentinel, blob, f"隐藏区内容值 {sentinel} 不得进入盘点")

    def test_collapsed_outline_group_counts_as_hidden(self):
        """折叠分组内的行视同隐藏（Excel 折叠后不可见）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "分组.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"], ws["A4"] = "头", 12345, 23456, "尾"
        ws.row_dimensions[2].outlineLevel = 1        # 组内行
        ws.row_dimensions[3].outlineLevel = 1        # 组内行
        ws.row_dimensions[4].collapsed = True        # 汇总行（默认在下方）= 组已折叠
        wb.save(f)

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/分组.xlsx", data_only=True)
        values = [c.value for row in rebuilt.active.iter_rows() for c in row if c.value is not None]
        self.assertEqual(["头", "尾"], values, "折叠组内的行必须排除")

    # ---- 3 同名消歧 ------------------------------------------------------
    def test_same_named_versions_do_not_overwrite_each_other(self):
        """定稿与送审稿同名 → 工作版两份都在，来源可区分。"""
        import openpyxl
        for stage, val in (("定稿", "出租"), ("送审稿", "自用")):
            d = self.src / stage
            d.mkdir()
            wb = openpyxl.Workbook()
            wb.active["A1"] = val
            wb.save(d / "13-评估明细表-20231231.xlsx")

        self._run()

        work = self.case / "工作版"
        names = sorted(p.name for p in work.glob("*.xlsx"))
        self.assertEqual(2, len(names), names)
        got = {
            openpyxl.load_workbook(work / n, data_only=True).active["A1"].value
            for n in names
        }
        self.assertEqual({"出租", "自用"}, got, "两份同名材料不得互相覆盖")
        records = self._inventory()
        self.assertEqual({"定稿", "送审稿"}, {r["stage"] for r in records})
        self.assertTrue(any(r.get("nameCollision") for r in records), "同名冲突须在盘点中标出")

    # ---- 4 源目录只读 ----------------------------------------------------
    def test_source_materials_directory_is_not_polluted(self):
        """运行后源材料目录的文件集合不得变化（不得落派生文件）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        wb.save(d / "a.xlsx")
        before = sorted(str(p.relative_to(self.src)) for p in self.src.rglob("*"))

        self._run()

        after = sorted(str(p.relative_to(self.src)) for p in self.src.rglob("*"))
        self.assertEqual(before, after)

    # ---- 4b 缺依赖不得静默空返 -------------------------------------------
    def test_docx_without_python_docx_never_silently_returns_empty(self):
        """缺 python-docx 时必须回退（或明确报因），不得静默产出空文本。"""
        import sys
        import zipfile as zf
        d = self.src / "定稿"
        d.mkdir()
        # 手工造最小 .docx（zip + 两个必需部件）
        f = d / "说明.docx"
        with zf.ZipFile(f, "w") as z:
            z.writestr("[Content_Types].xml",
                       '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
                       'package/2006/content-types"><Default Extension="xml" ContentType='
                       '"application/xml"/></Types>')
            z.writestr("word/document.xml",
                       '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
                       'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>评估说明正文</w:t>'
                       '</w:r></w:p></w:body></w:document>')

        saved = sys.modules.get("docx")
        sys.modules["docx"] = None          # 模拟未安装 python-docx
        try:
            self._run()
        finally:
            if saved is None:
                sys.modules.pop("docx", None)
            else:
                sys.modules["docx"] = saved

        rec = self._inventory()[0]
        if rec["readable"]:
            text = (self.case / "提取/说明.docx.txt").read_text(encoding="utf-8")
            self.assertTrue(text.strip(), "回退成功时不得产出空文本")
            self.assertIn("回退", rec.get("note", ""))
        else:
            self.assertTrue(rec.get("note"), "读不到必须写明原因，不得静默")

    # ---- 4c 隐藏区引用审计（只坐标，不读内容）-------------------------------
    def test_hidden_reference_audit_flags_visible_results_depending_on_hidden_inputs(self):
        """可见公式引用隐藏列/行/隐藏表 → 记**人工建议检查项**（不是 AI 问题）；且不得读隐藏内容。

        对应真实案件：土地表隐藏列 N（账面价值）被 32 处可见公式引用、底稿隐藏评分列
        被 24 处引用 —— 这类隐藏列**是计算输入**，一刀切剔除会让结果不可复核。
        去向（2026-09-16 用户口径）：审核只要可见区公式计算正确即可 → 只出建议项、不出 `issues[]`。
        """
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "引用审计.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "土地表"
        ws["P1"] = "=U25*2"                  # 相对引用隐藏列 U
        ws["P2"] = "=SUM($U$25:$U$30)"       # 绝对引用 + 范围（曾漏匹配 $U$25）
        ws["P3"] = "=A1*2"                   # 仅引用可见格 → 不得报
        ws["P4"] = "=LOG10(A1)"              # 函数名不得被误判为单元格引用
        ws["P5"] = "='隐藏表'!B2"            # 引用隐藏工作表
        ws["U25"] = 999999                   # 隐藏列的内容（哨兵）
        ws.column_dimensions["U"].hidden = True
        hid = wb.create_sheet("隐藏表")
        hid["B2"] = 888888
        hid.sheet_state = "hidden"
        wb.save(f)

        self._run()

        rec = self._inventory()[0]["workbook"]
        kinds = sorted({h["kind"] for h in rec["hiddenRefs"]})
        self.assertEqual(["引用隐藏列", "引用隐藏工作表"], kinds, rec["hiddenRefs"])
        cells = {h["cell"] for h in rec["hiddenRefs"]}
        self.assertEqual({"P1", "P2", "P5"}, cells, "只报真正引用隐藏区的可见格")
        self.assertEqual(["土地表!P1", "土地表!P2", "土地表!P5"], rec["calcChainNotReproducible"])
        self.assertIn("manualConfirmationItems", rec["hiddenRefDisposition"],
                      "去向必须随数据走：人工建议检查项，不作为 AI 问题")
        self.assertIn("不作为 AI 问题", rec["hiddenRefDisposition"])
        blob = (self.case / "材料盘点.json").read_text(encoding="utf-8")
        for sentinel in ("999999", "888888"):
            self.assertNotIn(sentinel, blob, "引用审计只出坐标，绝不能带出隐藏区内容")

    def test_hidden_reference_audit_is_quiet_when_hidden_columns_are_unreferenced(self):
        """隐藏列仅被隐藏区内部自引用（可见区无引用）→ 不产生建议项元数据（无声即无噪声）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "汇总表"
        ws["A1"] = 1
        ws["C1"] = "=C2*2"                   # 隐藏列内部自引用（隐藏格，不得读取）
        ws["C2"] = 5
        ws.column_dimensions["C"].hidden = True
        wb.save(d / "自引用.xlsx")

        self._run()

        rec = self._inventory()[0]["workbook"]
        self.assertEqual([], rec["hiddenRefs"])
        self.assertEqual([], rec["calcChainNotReproducible"])
        self.assertIsNone(rec["hiddenRefDisposition"], "无引用时不出任何建议项元数据")

    # ---- 4d2 多版本隐藏结构比对（H0 元数据级）----------------------------
    def test_same_named_versions_with_different_hidden_structure_flag_drift(self):
        """同名定稿/送审稿隐藏结构不一致 → 输出元数据级提示，且不读隐藏内容。"""
        import openpyxl
        for stage, hide_col in (("定稿", "O"), ("送审稿", None)):
            d = self.src / stage
            d.mkdir()
            wb = openpyxl.Workbook()
            wb.active["A1"] = "值"
            wb.active["O1"] = 999999
            if hide_col:
                wb.active.column_dimensions[hide_col].hidden = True
            wb.save(d / "13-评估明细表-20231231.xlsx")

        self._run()

        drift = self._payload().get("hiddenStructureDrift", [])
        self.assertEqual(1, len(drift), drift)
        variants = drift[0]["variants"]
        self.assertEqual(2, len(variants), variants)
        blob = (self.case / "材料盘点.json").read_text(encoding="utf-8")
        self.assertNotIn("999999", blob, "比对只出隐藏元数据，不得带出隐藏内容值")

    def test_same_named_versions_with_identical_hidden_structure_do_not_flag(self):
        import openpyxl
        for stage in ("定稿", "送审稿"):
            d = self.src / stage
            d.mkdir()
            wb = openpyxl.Workbook()
            wb.active["A1"] = "值"
            wb.active.column_dimensions["O"].hidden = True
            wb.save(d / "13-评估明细表-20231231.xlsx")

        self._run()

        self.assertEqual([], self._payload().get("hiddenStructureDrift", []))

    # ---- 4d 压缩包：下载后必须先解压再审核 --------------------------------
    def test_archive_is_extracted_and_contents_are_processed(self):
        """归档必须先解压，且解压产物随同一套逻辑继续处理（xlsx 也要出工作版）。"""
        import openpyxl
        import zipfile
        d = self.src / "参考材料"
        d.mkdir()
        inner = self.case / "tmp_inner.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "底稿"
        wb.save(inner)
        with zipfile.ZipFile(d / "测算表.zip", "w") as z:
            z.write(inner, "13-工业变性住宅-测算底稿.xlsx")
        inner.unlink()

        self._run()

        payload = self._payload()
        self.assertEqual(1, len(payload["archives"]))
        self.assertEqual(["13-工业变性住宅-测算底稿.xlsx"], payload["archives"][0]["extracted"])
        extracted = [r for r in self._inventory() if r.get("originArchive")]
        self.assertEqual(1, len(extracted), self._inventory())
        self.assertEqual("材料-源/参考材料/测算表.zip", extracted[0]["originArchive"])
        self.assertTrue(extracted[0]["workbook"]["workVersion"].startswith("工作版/"),
                        "解压出的 xlsx 必须同样重建工作版")

    def test_archive_with_gbk_filenames_is_decoded(self):
        """zip 内中文名常为 GBK（无 UTF-8 标志位）→ 必须正确解码，文件名可读。

        注意：`zipfile.writestr` 遇到非 ASCII 名会自动置 UTF-8 标志位，造不出真实 GBK 包；
        故用同长度 ASCII 占位写入后，在字节层替换为 GBK 字节（标志位保持 0）。
        """
        import zipfile
        d = self.src / "参考材料"
        d.mkdir()
        raw = "13-轻型汽车-土地使用权-工业变性住宅-测算底稿.xlsx".encode("gbk")
        placeholder = "A" * len(raw)
        path = d / "gbk.zip"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr(placeholder, b"x")
        blob = path.read_bytes()
        assert placeholder.encode("ascii") in blob
        path.write_bytes(blob.replace(placeholder.encode("ascii"), raw))
        with zipfile.ZipFile(path) as z:          # 自检：确实是"无 UTF-8 标志位"的 GBK 名
            info = z.infolist()[0]
            self.assertEqual(0, info.flag_bits & 0x800)
            self.assertEqual(raw.decode("cp437"), info.filename)   # cp437 读出的乱码名（真实情形）

        self._run()

        names = self._payload()["archives"][0]["extracted"]
        self.assertEqual(["13-轻型汽车-土地使用权-工业变性住宅-测算底稿.xlsx"], names)

    def test_archive_entries_escaping_or_symlinked_are_refused(self):
        """zip-slip / 绝对路径 / 符号链接一律拒绝，并如实记账。"""
        import zipfile
        d = self.src / "参考材料"
        d.mkdir()
        with zipfile.ZipFile(d / "evil.zip", "w") as z:
            z.writestr("ok.txt", "fine")
            z.writestr("../escape.txt", "bad")
            z.writestr("/abs/escape.txt", "bad")
            link = zipfile.ZipInfo("link")
            link.external_attr = (0o120777 << 16)
            z.writestr(link, "/etc/passwd")

        self._run()

        arec = self._payload()["archives"][0]
        self.assertEqual(["ok.txt"], arec["extracted"])
        self.assertEqual(3, len(arec["failures"]), arec["failures"])
        self.assertFalse((self.case / "escape.txt").exists())
        self.assertFalse((self.src / ".." / "escape.txt").resolve().exists())

    def test_archive_entries_that_are_review_records_are_not_extracted(self):
        """阶段一隔离门禁优先于解压：归档内复核记录类文件不解压。"""
        import zipfile
        d = self.src / "参考材料"
        d.mkdir()
        with zipfile.ZipFile(d / "mixed.zip", "w") as z:
            z.writestr("测算底稿.xlsx", "x")
            z.writestr("三级复核意见.docx", "secret")

        self._run()

        arec = self._payload()["archives"][0]
        self.assertEqual(["测算底稿.xlsx"], arec["extracted"])
        self.assertEqual(["三级复核意见.docx"], arec["skippedIsolation"])
        blob = (self.case / "材料盘点.json").read_text(encoding="utf-8")
        self.assertIn("三级复核意见.docx", blob)      # 只记名字
        self.assertNotIn("secret", blob)              # 不得带出内容

    def test_unsupported_archive_records_capability_gap_not_silence(self):
        """不支持的归档格式（.rar/.7z 且无外部工具）必须记 capability gap，不得静默跳过。"""
        import shutil as sh
        d = self.src / "参考材料"
        d.mkdir()
        (d / "扫描件.rar").write_bytes(b"Rar!\x1a\x07\x00dummy")
        saved = sh.which
        sh.which = lambda name: None
        try:
            self._run()
        finally:
            sh.which = saved

        arec = self._payload()["archives"][0]
        self.assertTrue(arec["capabilityGaps"], arec)
        self.assertIn("未安装", arec["capabilityGaps"][0])

    # ---- 5 格式语义 ------------------------------------------------------
    def test_number_format_is_preserved(self):
        """重建须保留 number_format，避免数值/日期语义被改变。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 0.075
        ws["A1"].number_format = "0.00%"
        wb.save(d / "格式.xlsx")

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/格式.xlsx")
        self.assertEqual("0.00%", rebuilt.active["A1"].number_format)


    # ---- 6 隐藏区区段展开（防"隐藏列漏剔"假阳性）------------------------
    @staticmethod
    def _hide_columns(path: Path, sheet_name: str, lo: int, hi: int):
        """按 raw XML 写 `<col min lo max hi hidden="1"/>`（含区段，覆盖假阳性根因场景）。

        注意：openpyxl 的 `ws.column_dimensions` 只把该区段挂在**首列**上，这正是此前
        `_hidden_metadata()` 漏剔区段内其余列、把人工隐藏内容读进工作版的根因。
        """
        import re
        tmp = path.with_suffix(".hide.xlsx")
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            names = dict(re.findall(
                r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"',
                zin.read("xl/workbook.xml").decode("utf-8")))
            rels = {}
            for rel in re.finditer(r"<Relationship\b([^>]*)/?>",
                                   zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")):
                a = dict(re.findall(r'([A-Za-z_:][\w:.-]*)\s*=\s*"([^"]*)"', rel.group(1)))
                if a.get("Id") and a.get("Target"):
                    rels[a["Id"]] = a["Target"]
            target = rels[names[sheet_name]].lstrip("/")
            target = target if target.startswith("xl/") else "xl/" + target
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == target:
                    text = data.decode("utf-8")
                    cols = '<cols><col min="{0}" max="{1}" hidden="1"/></cols>'.format(lo, hi)
                    text = re.sub(r"<sheetData", cols + "<sheetData", text, count=1)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp), str(path))

    def test_hidden_column_range_is_fully_excluded(self):
        """隐藏列**区段**（min<max）内每一列都不得进入工作版（此前只剔首列 → 假阳性）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "汇总表"
        ws["A1"] = "序号"
        ws["J1"] = "隐藏段首列"
        ws["K1"] = "隐藏段第二列(曾漏剔)"
        wb.save(d / "区段隐藏.xlsx")
        self._hide_columns(d / "区段隐藏.xlsx", "汇总表", lo=10, hi=11)   # J:K

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/区段隐藏.xlsx")
        rws = rebuilt["汇总表"]
        self.assertEqual("序号", rws["A1"].value)
        self.assertIsNone(rws["J1"].value, "区段内首列 J 应被剔除")
        self.assertIsNone(rws["K1"].value, "区段内其余列 K 也必须被剔除（本条即历史失效点）")
        meta = [i for i in self._inventory() if i["name"] == "区段隐藏.xlsx"][0]["workbook"]
        self.assertIn("K", meta["hiddenMeta"]["hiddenCols"]["汇总表"])

    # ---- 7 表名含首尾空格时隐藏区引用审计仍须命中 ------------------------
    def test_hidden_reference_audit_matches_sheet_name_with_spaces(self):
        """表名含首尾空格时，隐藏**列**引用仍须被审计命中（此前恒返回 0 处 → 漏报）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "报价表  "          # 尾空格，历史失效场景
        ws["A1"] = 1
        ws["B1"] = "=A1*2"
        wb.save(d / "空格表名.xlsx")
        self._hide_columns(d / "空格表名.xlsx", "报价表  ", lo=1, hi=1)   # 隐藏 A 列

        self._run()

        meta = [i for i in self._inventory() if i["name"] == "空格表名.xlsx"][0]["workbook"]
        self.assertGreaterEqual(meta["hiddenRefsCount"], 1,
                                "表名含空格时不得漏报隐藏列引用（归一化缺失的回归）")
        self.assertTrue(any(r["kind"] == "引用隐藏列" for r in meta["hiddenRefs"]), meta["hiddenRefs"])

    # ---- 8 虚增 dimension 不得拖垮遍历 ----------------------------------
    def test_inflated_dimension_does_not_force_full_sheet_scan(self):
        """`dimension` 被虚增（如 A1:Y1048575）时，遍历边界须取**真实用区**。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "合计"
        wb.save(d / "虚增维度.xlsx")
        p = d / "虚增维度.xlsx"
        tmp = p.with_suffix(".dim.xlsx")
        with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet"):
                    text = data.decode("utf-8")
                    text = text.replace('ref="A1"', 'ref="A1:Y1048575"', 1)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp), str(p))

        src_ws = openpyxl.load_workbook(p).active
        max_row, max_col = self.module._sheet_bounds(src_ws)
        self.assertEqual((1, 1), (max_row, max_col), "虚增 dimension 不得被当作遍历边界")

        self._run()      # 仍须在正常时间内完成
        meta = [i for i in self._inventory() if i["name"] == "虚增维度.xlsx"][0]["workbook"]
        self.assertEqual(1, meta["sheets"][0]["maxRow"])


    # ---- 9 表格遍历的四层边界（B1 扫描 / B2 内容 / B3 异常 / B4 护栏） ------
    def test_orphan_styled_cell_does_not_force_full_sheet_scan(self):
        """游离格式格（`value=None` 但有样式）落在末行：边界须按**有值格**定，并登记规范提示。

        真实失效（2026-302135-LX9619-BG8634 的 `4-15-3无形-其他`）：
        944 个存在格中 697 个纯格式 + 1 个游离在第 1048575 行
        → 重建循环 26,214,375 次、单表 56 s；隐藏引用审计再扫满 → 70 s。
        只按"存在的格"（`_cells`）定界**仍会踩到**——那个游离格就在 `_cells` 里。
        """
        import openpyxl
        from openpyxl.styles import Font
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "合计"
        ws["A2"] = 42
        ws.cell(row=1048575, column=6).font = Font(bold=True)      # 只有样式、没有值
        wb.save(d / "游离格式格.xlsx")

        src_ws = openpyxl.load_workbook(d / "游离格式格.xlsx").active
        self.assertEqual((2, 1), self.module._sheet_bounds(src_ws),
                         "边界须按有值格定，不得被游离格式格顶到 1048575 行")

        self._run()                                                 # 且须在正常时间内完成

        meta = [i for i in self._inventory() if i["name"] == "游离格式格.xlsx"][0]["workbook"]
        self.assertEqual(2, meta["sheets"][0]["maxRow"])
        anomalies = [a for a in self._payload().get("sheetAnomalies", [])
                     if a["path"].endswith("游离格式格.xlsx")]
        self.assertEqual(1, len(anomalies), "声明用区远大于实际有值区时须登记表格规范提示")
        self.assertGreater(anomalies[0]["orphanRows"], 1000)
        self.assertIn("1048575", anomalies[0]["declaredDim"])
        self.assertEqual(1, anomalies[0]["styleOnlyCells"])
        self.assertEqual({"maxRow": 1048575, "maxCol": 6}, meta["sheets"][0]["declaredSpan"],
                         "声明用区与有值区不一致时须一并留痕（低于 B3 阈值时更要有）")
        rebuilt = openpyxl.load_workbook(self.case / "工作版/游离格式格.xlsx", data_only=True)
        self.assertEqual("合计", rebuilt.active["A1"].value)
        self.assertEqual(42, rebuilt.active["A2"].value)

    def test_oversized_sheet_records_capability_gap_not_silence(self):
        """单表有值格超上限：记 capability gap 并跳过该表，其余表继续，不得静默。"""
        import openpyxl
        orig = self.module.MAX_SHEET_CELLS
        self.module.MAX_SHEET_CELLS = 5
        try:
            d = self.src / "定稿"
            d.mkdir()
            wb = openpyxl.Workbook()
            big = wb.active
            big.title = "超限表"
            for i in range(1, 11):
                big.cell(row=i, column=1, value=i)
            wb.create_sheet("正常表")["A1"] = "ok"
            wb.save(d / "超限.xlsx")
            self._run()
        finally:
            self.module.MAX_SHEET_CELLS = orig

        rec = [i for i in self._inventory() if i["name"] == "超限.xlsx"][0]
        self.assertTrue(rec.get("capabilityGaps"), "超限必须记 capability gap，不得静默跳过")
        self.assertIn("未完整处理", " ".join(rec["capabilityGaps"]))
        self.assertEqual(["正常表"], [s["sheet"] for s in rec["workbook"]["sheets"]],
                         "超限表不进工作版，其余表照常处理")

    # ---- 10 非单元格证据（媒体）不得因工作版丢失而被判"缺失" ----------------
    def test_raw_media_is_counted_and_flagged_as_non_cell_evidence(self):
        """原件含内嵌媒体时：登记 `rawMediaCount`，汇总进 `nonCellEvidence` 并提示不得据工作版判缺失。

        真实失效（2026-302150-LX9757-BG8677）：重建工作版丢失全部 `xl/media`（审核对象已非送审件原貌），
        导致 `MKT-004` 判"可比实例位置图为空、询价截图缺失"，而原件实有 26 个媒体对象。
        """
        import openpyxl
        import shutil
        import zipfile
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        p = d / "带图.xlsx"
        wb.save(p)
        tmp = p.with_suffix(".media.xlsx")          # openpyxl 不写图，直接补 zip 部件模拟内嵌图
        with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("xl/media/image1.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64)
        shutil.move(str(tmp), str(p))

        self._run()

        rec = [i for i in self._inventory() if i["name"] == "带图.xlsx"][0]["workbook"]
        self.assertEqual(1, rec["rawMediaCount"], "须登记原件媒体数")
        self.assertEqual(0, rec["mediaCarriedOver"], "重建法不复制媒体，须如实记账")
        self.assertIn("禁止依据工作版", rec["nonCellEvidenceNote"])
        nce = self._payload()["nonCellEvidence"]
        self.assertEqual((1, 1), (nce["filesWithMedia"], nce["rawMediaCount"]))
        rebuilt = self.case / "工作版/带图.xlsx"
        self.assertEqual(
            [], [n for n in zipfile.ZipFile(rebuilt).namelist() if n.startswith("xl/media/")],
            "工作版不得残留媒体——正因如此，媒体类结论必须回 raw 原件")

    # ---- 11 媒体证据通道：docx 与独立图片也必须导出（此前只有 xlsx 计数）----
    def test_docx_and_standalone_images_are_exported_into_media_index(self):
        """docx 内嵌图与独立图片必须落 `媒体证据/` 并进 `媒体索引.json`。

        真实失效：`评估说明.docx` 正文写着"以下为询价截图："，图片被文本抽取静默丢弃，
        盘点里连媒体字段都没有；独立图片的 note 只有"图片：无 OCR"，没有任何可读路径。
        """
        d = self.src / "定稿"
        d.mkdir()
        _minimal_docx_with_image(d / "评估说明.docx")
        (d / "现场照片.png").write_bytes(_PNG)

        self._run()

        items = {r["name"]: r for r in self._inventory()}
        docx_ev = items["评估说明.docx"]["mediaEvidence"]
        self.assertEqual(1, docx_ev["exportedCount"], docx_ev)
        self.assertEqual(1, items["现场照片.png"]["mediaEvidence"]["exportedCount"])
        self.assertEqual("host-vision", items["现场照片.png"]["mediaEvidence"]["evidenceChannel"])

        idx = json.loads((self.case / "媒体索引.json").read_text(encoding="utf-8"))
        self.assertEqual(2, idx["count"], idx)
        for e in idx["entries"]:
            self.assertTrue((self.case / e["localPath"]).exists(), e)
            self.assertTrue(e["sha256"] and e["sizeBytes"] > 0, e)
        locators = sorted(e["locator"] for e in idx["entries"])
        self.assertTrue(any("图片#" in x for x in locators), locators)
        self.assertTrue(any("（整图）" in x for x in locators), locators)

        nce = self._payload()["nonCellEvidence"]
        self.assertEqual(2, nce["filesWithMedia"])
        self.assertEqual(2, nce["exportedMediaCount"])
        self.assertEqual("媒体索引.json", nce["mediaIndexPath"])

    def test_unreferenced_media_is_not_silently_dropped(self):
        """未被 drawing 引用的媒体部件：不导出，但必须留未核原因（未核 ≠ 缺失）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        p = d / "无锚点图.xlsx"
        wb.save(p)
        tmp = p.with_suffix(".media.xlsx")
        with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("xl/media/image1.png", _PNG)
        shutil.move(str(tmp), str(p))

        self._run()

        rec = [r for r in self._inventory() if r["name"] == "无锚点图.xlsx"][0]
        ev = rec["mediaEvidence"]
        self.assertEqual(0, ev["exportedCount"], "锚点不可得时不得导出")
        self.assertTrue(ev["unresolvedCount"] >= 1, ev)
        self.assertIn("未导出", " ".join(ev["unresolvedReasons"]))
        self.assertFalse((self.case / "媒体证据").exists() and
                         any((self.case / "媒体证据").rglob("*.png")),
                         "锚点不可得的媒体不得落盘")

    # ---- 12 阶段二复跑不得覆盖阶段一冻结产物 ------------------------------
    def test_review_pass_does_not_clobber_phase1_artifacts(self):
        """阶段二对 `复核-人工/` 复跑：抽复核件媒体，但阶段一冻结产物逐字节不变。"""
        d = self.src / "定稿"
        d.mkdir()
        _minimal_docx_with_image(d / "评估说明.docx")
        self._run()
        inv_before = _sha256(self.case / "材料盘点.json")
        idx_before = _sha256(self.case / "媒体索引.json")

        review = self.case / "复核-人工"
        review.mkdir()
        _minimal_docx_with_image(review / "三级复核意见.docx")

        self._run_label("复核", review)

        self.assertEqual(inv_before, _sha256(self.case / "材料盘点.json"),
                         "阶段一盘点表是冻结产物，阶段二复跑不得改写")
        self.assertEqual(idx_before, _sha256(self.case / "媒体索引.json"),
                         "阶段一媒体索引同样不得被阶段二覆盖")
        idx = json.loads((self.case / "复核媒体索引.json").read_text(encoding="utf-8"))
        self.assertEqual(1, idx["count"], idx)
        self.assertIn("复核意见", json.dumps(idx, ensure_ascii=False))
        self.assertTrue((self.case / "复核盘点.json").exists())
        # 复核件的媒体落在独立的媒体证据目录，不混入阶段一目录
        self.assertTrue(next((self.case / "媒体证据-复核").rglob("*.png"), None))


    # ---- 13 CLI：相对 --src 必须相对 --case 解析 --------------------------
    def test_cli_relative_src_resolves_against_case(self):
        """`--src 复核-人工` 必须解析到 `<案例>/复核-人工`，而不是当前工作目录。

        真实踩坑：SKILL.md 里的阶段二命令是 `--case <案例目录> --src 复核-人工 --label 复核`，
        若相对路径按 cwd 解析，命令会直接报"源材料目录不存在"。
        """
        import sys
        review = self.case / "复核-人工"
        review.mkdir()
        _minimal_docx_with_image(review / "三级复核意见.docx")

        saved = sys.argv
        sys.argv = ["prepare_materials.py", "--case", str(self.case),
                    "--src", "复核-人工", "--label", "复核"]
        try:
            rc = self.module.main()
        finally:
            sys.argv = saved

        self.assertEqual(0, rc)
        idx = json.loads((self.case / "复核媒体索引.json").read_text(encoding="utf-8"))
        self.assertEqual(1, idx["count"], idx)
        self.assertFalse((self.case / "材料盘点.json").exists(),
                         "只跑阶段二时不应产出阶段一盘点表")


if __name__ == "__main__":
    unittest.main(verbosity=2)
